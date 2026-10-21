"""Tests for document extraction (PDF / DOCX / TXT), validation and segmentation."""
import io
from pathlib import Path

import docx
import pymupdf
import pytest

from nlp.document_processor import (
    CorruptedDocumentError,
    DocumentError,
    EmptyDocumentError,
    FileTooLargeError,
    NoExtractableTextError,
    UnsupportedFileTypeError,
    extract_document,
    extract_document_from_path,
)

SAMPLES = Path(__file__).resolve().parent.parent / "sample_data"


# ------------------------------------------------------------ file builders
def make_pdf(pages):
    pdf = pymupdf.open()
    for lines in pages:
        page = pdf.new_page()
        y = 80
        for line in lines:
            page.insert_text((72, y), line, fontsize=11)
            y += 40
    data = pdf.tobytes()
    pdf.close()
    return data


def make_docx(paragraphs, table=None):
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------- PDF
def test_pdf_extraction_with_page_references():
    data = make_pdf([["The exam starts at 10 AM sharp.", "Students must carry identity cards."],
                     ["The library is closed on Sundays."]])
    doc = extract_document("rules.pdf", data)
    assert doc.file_type == "pdf" and doc.page_count == 2
    assert doc.num_sentences == 3
    last = doc.sentences[-1]
    assert last.text == "The library is closed on Sundays."
    assert last.page == 2 and last.paragraph == 1
    assert last.source == "rules.pdf"
    assert "page 2" in last.location and "rules.pdf" in last.location


def test_pdf_without_text_is_reported_not_crashed():
    pdf = pymupdf.open(); pdf.new_page(); data = pdf.tobytes(); pdf.close()
    with pytest.raises(NoExtractableTextError) as exc:
        extract_document("scan.pdf", data)
    assert "OCR" in str(exc.value)


def test_corrupted_pdf():
    with pytest.raises(CorruptedDocumentError):
        extract_document("broken.pdf", b"this is definitely not a pdf file")


def test_sample_pdf_extracts():
    doc = extract_document_from_path(str(SAMPLES / "circular_a_original_policy.pdf"))
    assert doc.page_count == 2
    assert any("75%" in s.text for s in doc.sentences)
    assert any(s.page == 2 for s in doc.sentences)


# -------------------------------------------------------------------- DOCX
def test_docx_extraction_paragraph_references_and_tables():
    data = make_docx(["The fee is due in July.", "Hostel rules apply to all students."],
                     table=[["Late fee applies after the due date.", "Fine is fixed."]])
    doc = extract_document("notice.docx", data)
    texts = [s.text for s in doc.sentences]
    assert "The fee is due in July." in texts
    assert "Late fee applies after the due date." in texts   # table cell text is read
    first = doc.sentences[0]
    assert first.paragraph == 1 and first.page is None
    assert "paragraph 1" in first.location


def test_corrupted_docx():
    with pytest.raises(CorruptedDocumentError):
        extract_document("broken.docx", b"PK\x03\x04 not really a docx")


def test_empty_docx():
    with pytest.raises(EmptyDocumentError):
        extract_document("empty.docx", make_docx([]))


def test_sample_docx_extracts():
    doc = extract_document_from_path(str(SAMPLES / "circular_b_revised_notice.docx"))
    assert any("65%" in s.text for s in doc.sentences)


# --------------------------------------------------------------------- TXT
def test_txt_extraction_paragraphs():
    data = b"First paragraph is here. It has two sentences.\n\nSecond paragraph is here."
    doc = extract_document("notes.txt", data)
    assert [s.paragraph for s in doc.sentences] == [1, 1, 2]
    assert doc.num_paragraphs == 2


@pytest.mark.parametrize("encoded", [
    "Caf\u00e9 is open on weekdays only.".encode("utf-8"),
    "Caf\u00e9 is open on weekdays only.".encode("cp1252"),
    "Caf\u00e9 is open on weekdays only.".encode("utf-16"),
])
def test_txt_encodings(encoded):
    doc = extract_document("menu.txt", encoded)
    assert doc.sentences[0].text.endswith("open on weekdays only.")


def test_txt_binary_content_rejected():
    with pytest.raises(CorruptedDocumentError):
        extract_document("fake.txt", b"abc\x00\x01\x02def")


def test_sample_txt_files_extract():
    for name in ("circular_a_original_policy.txt", "circular_b_revised_notice.txt"):
        assert extract_document_from_path(str(SAMPLES / name)).num_sentences >= 10


# ------------------------------------------------------ validation / errors
@pytest.mark.parametrize("name", ["virus.exe", "old.doc", "noextension", "image.png", "data.csv"])
def test_unsupported_file_types(name):
    with pytest.raises(UnsupportedFileTypeError):
        extract_document(name, b"some bytes")


def test_zero_byte_and_whitespace_only_files():
    with pytest.raises(EmptyDocumentError):
        extract_document("empty.txt", b"")
    with pytest.raises(EmptyDocumentError):
        extract_document("blank.txt", b"   \n\n \t \n")


def test_oversized_file_rejected():
    with pytest.raises(FileTooLargeError):
        extract_document("big.txt", b"a" * 100, max_bytes=50)


def test_all_errors_share_base_class():
    for exc in (UnsupportedFileTypeError, FileTooLargeError, EmptyDocumentError,
                CorruptedDocumentError, NoExtractableTextError):
        assert issubclass(exc, DocumentError)


# ------------------------------------------------------------- segmentation
def test_short_fragments_are_dropped_but_not_if_nothing_else_remains():
    doc = extract_document("a.txt", b"Page 1\n\nThis is a proper full sentence here.")
    assert [s.text for s in doc.sentences] == ["This is a proper full sentence here."]
    only_short = extract_document("b.txt", b"Yes.\n\nNo.")
    assert only_short.num_sentences == 2


def test_sentence_indexes_are_sequential_and_context_kept():
    doc = extract_document("a.txt", b"One sentence is here. Another sentence is here.")
    assert [s.index for s in doc.sentences] == [0, 1]
    assert all(s.context.startswith("One sentence") for s in doc.sentences)
