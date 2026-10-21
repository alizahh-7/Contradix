"""Document text extraction + sentence segmentation (PDF / DOCX / TXT).

Pipeline step:  uploaded bytes -> paragraphs (with page/paragraph refs)
                -> sentences (SentenceRecord, each remembering where it came from).

* PDF  : PyMuPDF (text PDFs only - scanned/image PDFs need OCR, which is out of scope)
* DOCX : python-docx (paragraphs and table cells, in document order)
* TXT  : standard library decoding

Heavy libraries are imported lazily so unit tests / the rest of the app work even
if a library is missing (a clear error is raised only when that format is used).
"""
from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass, field

from nlp.preprocessing import clean_text, split_sentences_batch

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")
MAX_FILE_MB = 10
MAX_FILE_BYTES = MAX_FILE_MB * 1024 * 1024
MAX_SENTENCES_PER_DOC = 3000
MIN_WORDS = 3  # shorter fragments (headings, page numbers) are dropped


# ------------------------------------------------------------------- errors
class DocumentError(Exception):
    """Base class: the message is always safe to show to the user."""


class UnsupportedFileTypeError(DocumentError):
    pass


class FileTooLargeError(DocumentError):
    pass


class EmptyDocumentError(DocumentError):
    pass


class CorruptedDocumentError(DocumentError):
    pass


class NoExtractableTextError(DocumentError):
    pass


# --------------------------------------------------------------- data classes
@dataclass
class Paragraph:
    text: str
    page: int | None = None      # PDF only
    index: int = 0               # 1-based paragraph number (page-local for PDF)


@dataclass
class SentenceRecord:
    text: str
    source: str                  # file name, or "Entered statement"
    page: int | None = None
    paragraph: int | None = None
    index: int = 0               # 0-based position among the document's sentences
    context: str = ""            # the full paragraph, used to highlight evidence

    @property
    def location(self) -> str:
        parts = [self.source]
        if self.page is not None:
            parts.append(f"page {self.page}")
        if self.paragraph is not None:
            parts.append(f"paragraph {self.paragraph}")
        return ", ".join(parts)


@dataclass
class ExtractedDocument:
    name: str
    file_type: str
    paragraphs: list[Paragraph] = field(default_factory=list)
    sentences: list[SentenceRecord] = field(default_factory=list)
    page_count: int | None = None
    truncated: bool = False

    @property
    def num_sentences(self) -> int:
        return len(self.sentences)

    @property
    def num_paragraphs(self) -> int:
        return len(self.paragraphs)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.paragraphs)


# ----------------------------------------------------------------- extractors
def _extract_pdf(data: bytes, name: str):
    try:
        import pymupdf
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("PDF support needs PyMuPDF. Run: pip install pymupdf") from exc
    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise CorruptedDocumentError(f"'{name}' could not be opened - the PDF looks corrupted.") from exc
    try:
        if pdf.page_count == 0:
            raise CorruptedDocumentError(f"'{name}' could not be read - the PDF has no pages or is corrupted.")
        if pdf.needs_pass:
            raise DocumentError(f"'{name}' is password-protected. Please upload an unlocked copy.")
        paragraphs: list[Paragraph] = []
        for page_no, page in enumerate(pdf, start=1):
            try:
                blocks = page.get_text("blocks", sort=True)
            except Exception as exc:
                raise CorruptedDocumentError(f"'{name}': page {page_no} could not be read.") from exc
            idx = 0
            for block in blocks:
                if len(block) > 6 and block[6] != 0:  # skip image blocks
                    continue
                text = clean_text(str(block[4]))
                if text:
                    idx += 1
                    paragraphs.append(Paragraph(text=text, page=page_no, index=idx))
        return paragraphs, pdf.page_count
    finally:
        pdf.close()


def _extract_docx(data: bytes, name: str):
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph as DocxParagraph
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("DOCX support needs python-docx. Run: pip install python-docx") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise CorruptedDocumentError(f"'{name}' could not be opened - the DOCX looks corrupted.") from exc

    texts: list[str] = []
    try:
        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                texts.append(DocxParagraph(child, document).text)
            elif tag == "tbl":
                for row in Table(child, document).rows:
                    seen = set()
                    for cell in row.cells:
                        if id(cell._tc) in seen:  # merged cells repeat
                            continue
                        seen.add(id(cell._tc))
                        texts.append(cell.text)
    except Exception as exc:
        raise CorruptedDocumentError(f"'{name}' could not be read - the DOCX looks corrupted.") from exc

    paragraphs, idx = [], 0
    for t in texts:
        t = clean_text(t)
        if t:
            idx += 1
            paragraphs.append(Paragraph(text=t, page=None, index=idx))
    return paragraphs, None


def _decode_text(data: bytes, name: str) -> str:
    is_utf16 = data.startswith((b"\xff\xfe", b"\xfe\xff"))
    if b"\x00" in data[:4096] and not is_utf16:
        raise CorruptedDocumentError(f"'{name}' does not look like a plain-text file.")
    # utf-16 is only tried when a BOM is present (otherwise it "decodes" any even-length bytes as garbage)
    encodings = ("utf-16",) if is_utf16 else ("utf-8-sig", "cp1252", "latin-1")
    for enc in encodings:
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    raise CorruptedDocumentError(f"'{name}' could not be decoded as text.")


def _extract_txt(data: bytes, name: str):
    text = _decode_text(data, name)
    blocks = re.split(r"\n\s*\n", text.replace("\r\n", "\n").replace("\r", "\n"))
    paragraphs, idx = [], 0
    for b in blocks:
        b = clean_text(b)
        if b:
            idx += 1
            paragraphs.append(Paragraph(text=b, page=None, index=idx))
    return paragraphs, None


_EXTRACTORS = {".pdf": _extract_pdf, ".docx": _extract_docx, ".txt": _extract_txt}


# ------------------------------------------------------------- segmentation
def _word_count(text: str) -> int:
    return len(re.findall(r"\w+", text))


def segment_paragraphs(paragraphs: list[Paragraph], source: str,
                       min_words: int = MIN_WORDS,
                       max_sentences: int = MAX_SENTENCES_PER_DOC):
    """Split paragraphs into SentenceRecords. Returns (records, truncated_flag)."""
    split = split_sentences_batch([p.text for p in paragraphs])
    all_records: list[SentenceRecord] = []
    for para, sents in zip(paragraphs, split):
        for s in sents:
            all_records.append(SentenceRecord(
                text=s, source=source, page=para.page, paragraph=para.index, context=para.text))
    kept = [r for r in all_records if _word_count(r.text) >= min_words]
    if not kept:          # e.g. a document made only of very short lines
        kept = all_records
    truncated = len(kept) > max_sentences
    kept = kept[:max_sentences]
    for i, r in enumerate(kept):
        r.index = i
    return kept, truncated


# --------------------------------------------------------------- public API
def get_extension(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def extract_document(filename: str, data: bytes, max_bytes: int = MAX_FILE_BYTES,
                     min_words: int = MIN_WORDS) -> ExtractedDocument:
    """Validate + extract + segment one uploaded document.

    Raises a DocumentError subclass (message is user-friendly) for unsupported
    types, empty / oversized / corrupted files and files with no extractable text.
    """
    ext = get_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext or filename}'. Supported formats: "
            + ", ".join(e.upper().lstrip('.') for e in SUPPORTED_EXTENSIONS) + "."
        )
    if not data:
        raise EmptyDocumentError(f"'{filename}' is empty (0 bytes).")
    if len(data) > max_bytes:
        raise FileTooLargeError(
            f"'{filename}' is {len(data) / 1024 / 1024:.1f} MB; the limit is {max_bytes / 1024 / 1024:.0f} MB."
        )

    paragraphs, page_count = _EXTRACTORS[ext](data, filename)
    if not paragraphs:
        if ext == ".pdf":
            raise NoExtractableTextError(
                f"No extractable text found in '{filename}'. It may be a scanned or image-only PDF; "
                "OCR is not supported in this project."
            )
        raise EmptyDocumentError(f"'{filename}' contains no readable text.")

    sentences, truncated = segment_paragraphs(paragraphs, filename, min_words=min_words)
    if not sentences:
        raise NoExtractableTextError(f"No sentences could be extracted from '{filename}'.")
    return ExtractedDocument(name=filename, file_type=ext.lstrip("."), paragraphs=paragraphs,
                             sentences=sentences, page_count=page_count, truncated=truncated)


def extract_document_from_path(path: str, **kwargs) -> ExtractedDocument:
    """Convenience wrapper: read a file from disk and extract it."""
    with open(path, "rb") as fh:
        data = fh.read()
    return extract_document(os.path.basename(path), data, **kwargs)
